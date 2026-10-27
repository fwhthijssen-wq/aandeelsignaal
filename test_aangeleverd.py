"""Aangeleverde jaarcijfers (aanvulsjabloon): inlezen en gaten vullen zonder officiële cijfers te raken."""
import copy, sys, tempfile, unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent / 'hulp'))
import aangeleverd_inlezen as AI
import bouw_webdata as B

GEEL = PatternFill('solid', fgColor='FFFF00')
GRIJS = PatternFill('solid', fgColor='E7E6E6')
KOP = ['ticker', 'naam', 'year', 'end', 'currency'] + AI.VELDEN + ['controle_eps', 'Opmerking']


def sjabloon(rijen):
    """rijen: (ticker, jaar, valuta, {veld: (waarde, 'geel'|'grijs')}, opmerking)"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Aanvullen'
    ws.append([k.title() for k in KOP]); ws.append(KOP)
    for t, j, v, cellen, opm in rijen:
        ws.append([t, t, str(j), f'{j}-12-31', v] + [cellen.get(k, (None,))[0] for k in AI.VELDEN] + [None, opm])
        r = ws.max_row
        for i, k in enumerate(AI.VELDEN):
            ws.cell(r, 6 + i).fill = GEEL if cellen.get(k, (None, 'geel'))[1] == 'geel' else GRIJS
    h = wb.create_sheet('Herkomst'); h.append(['Ticker', 'Boekjaar', 'Veld', 'Bron', 'Doc', 'Datum', 'Pagina', 'Opmerking']); h.append([None] * 8)
    for t, j, v, cellen, opm in rijen:
        h.append([t, j, 'alle', 'database (eodhd)', 'x', None, None, 'Valuta ' + v])
    pad = Path(tempfile.mkdtemp()) / 'in.xlsx'; wb.save(pad)
    return pad


class Inlezen(unittest.TestCase):
    def test_alleen_gele_cellen_eps_controle_en_pence(self):
        pad = sjabloon([
            ('A.AS', 2018, 'EUR', {'comparable_eps': (2.0, 'geel'), 'net_income': (200, 'geel'), 'shares_diluted': (100, 'geel'), 'sales': (5000, 'grijs')}, None),
            ('A.AS', 2019, 'EUR', {'comparable_eps': (9.0, 'geel'), 'net_income': (200, 'grijs'), 'shares_diluted': (100, 'grijs'), 'sales': (-5, 'geel')}, 'omzet klopt niet'),
            ('B.L', 2020, 'GBP', {'comparable_eps': (458.0, 'geel'), 'net_income': (458, 'geel'), 'shares_diluted': (100, 'geel')}, None),
        ])
        data, log, corr = AI.lees(pad)
        self.assertEqual(data['A.AS']['2018']['comparable_eps'], 2.0)
        self.assertNotIn('sales', data['A.AS']['2018'])                 # grijs: nooit overnemen
        self.assertNotIn('2019', data['A.AS'])                          # eps wijkt af (9 vs 2) en omzet negatief: niets over
        self.assertEqual(data['B.L']['2020']['comparable_eps'], 4.58)   # pence -> pond
        self.assertEqual(data['A.AS']['2018']['bron'], 'EODHD')
        self.assertEqual(log['eps_wijkt_af'], 1); self.assertEqual(log['onmogelijk'], 1); self.assertEqual(log['pence'], 1)
        self.assertEqual(corr[0]['opmerking'], 'omzet klopt niet')


def reg(jaren):
    return {'bronnen': {'EODHD': 1}, 'fondsen': {'X': {str(j): r for j, r in jaren.items()}}}


class Vullen(unittest.TestCase):
    def g(self, eps_jaren=(2021, 2022, 2023, 2024, 2025)):
        return {'niveau': 'jaarverslag', 'currencies': ['EUR'], 'jaren': [{'year': y, 'end': f'{y}-12-31', 'currency': 'EUR', 'comparable_eps': 1.0 + (y - 2021) * 0.1, 'sales': 1000.0} for y in eps_jaren]}

    def test_vult_gaten_nooit_over_officieel_en_wordt_indicatief(self):
        g = self.g()
        B.vul_aangeleverd(g, 'X', reg({2020: {'comparable_eps': 0.9, 'sales': 950, 'currency': 'EUR', 'end': '2020-12-31'},
                                       2021: {'comparable_eps': 5.0, 'sales': 9, 'net_income': 80, 'currency': 'EUR'}}))
        j = {x['year']: x for x in g['jaren']}
        self.assertEqual(j[2020]['comparable_eps'], 0.9); self.assertEqual(j[2020]['eps_bron'], 'aangeleverd')
        self.assertEqual(j[2021]['comparable_eps'], 1.0); self.assertEqual(j[2021]['sales'], 1000.0)   # officieel blijft
        self.assertEqual(j[2021]['net_income'], 80)                                                       # gat gevuld
        self.assertEqual(g['niveau'], 'indicatief')
        self.assertEqual(g['aangeleverd']['eps_jaren'], 1)

    def test_eenheid_valuta_en_lange_reeks(self):
        g = self.g()
        B.vul_aangeleverd(g, 'X', reg({2020: {'sales': 950000, 'currency': 'EUR'}, 2019: {'comparable_eps': 0.8, 'currency': 'USD'}}))
        self.assertNotIn(2019, [x['year'] for x in g['jaren']])
        self.assertNotIn(2020, [x['year'] for x in g['jaren']])                 # factor 950: andere eenheid, geweigerd
        self.assertEqual(g['aangeleverd']['valuta_jaren'], [{'jaar': 2019, 'valuta': 'USD', 'nodig': 'EUR'}])
        self.assertEqual(g['aangeleverd']['afgewezen'][0]['veld'], 'sales'); self.assertEqual(g['aangeleverd']['velden'], {})
        self.assertEqual(g['niveau'], 'jaarverslag')
        lang = self.g(range(2016, 2022))                                        # 6 jaar eigen winst per aandeel
        B.vul_aangeleverd(lang, 'X', reg({2015: {'comparable_eps': 0.9, 'sales': 990, 'currency': 'EUR'}}))
        r = [x for x in lang['jaren'] if x['year'] == 2015][0]
        self.assertNotIn('comparable_eps', r); self.assertEqual(r['sales'], 990)
        self.assertEqual(lang['niveau'], 'jaarverslag')

    def test_opnieuw_toepassen_en_nieuw_bestand(self):
        g = self.g()
        r1 = reg({2020: {'comparable_eps': 0.9, 'currency': 'EUR'}})
        B.vul_aangeleverd(g, 'X', r1); een = copy.deepcopy(g)
        B.vul_aangeleverd(g, 'X', r1)
        self.assertEqual(g, een)
        B.vul_aangeleverd(g, 'X', {'fondsen': {}})                              # bestand zonder dit fonds: alles terug
        self.assertNotIn(2020, [x['year'] for x in g['jaren']])
        self.assertEqual(g['niveau'], 'jaarverslag'); self.assertNotIn('aangeleverd', g)

    def test_fonds_zonder_jaarcijfers(self):
        g = {'fout': 'geen jaarcijfers', 'jaren': []}
        B.vul_aangeleverd(g, 'X', reg({y: {'comparable_eps': 1 + y / 1000, 'sales': 100, 'currency': 'EUR'} for y in range(2018, 2022)}))
        self.assertNotIn('fout', g); self.assertEqual(len(g['jaren']), 4); self.assertEqual(g['niveau'], 'indicatief')
        B.vul_aangeleverd(g, 'X', {'fondsen': {}})
        self.assertEqual(g['fout'], 'geen jaarcijfers'); self.assertEqual(g['jaren'], [])
        g2 = {'fout': 'geen', 'jaren': []}
        B.vul_aangeleverd(g2, 'X', reg({2020: {'comparable_eps': 1, 'currency': 'EUR'}}))   # te weinig jaren
        self.assertEqual(g2['fout'], 'geen')


if __name__ == '__main__':
    unittest.main()


class TweedeRonde(unittest.TestCase):
    def test_samenvoegen_nieuw_wint_en_andere_valuta_vervangt_rij(self):
        oud = {'X': {'2019': {'sales': 1, 'net_income': 2, 'currency': 'EUR'}, '2020': {'sales': 5, 'currency': 'EUR'}}}
        nieuw = {'X': {'2019': {'sales': 10, 'currency': 'EUR'}, '2020': {'sales': 7, 'currency': 'USD'}}, 'Y': {'2021': {'sales': 3, 'currency': 'GBP'}}}
        m = AI.samenvoegen(oud, nieuw)
        self.assertEqual(m['X']['2019'], {'sales': 10, 'net_income': 2, 'currency': 'EUR'})
        self.assertEqual(m['X']['2020'], {'sales': 7, 'currency': 'USD'})
        self.assertIn('Y', m)
        self.assertEqual(oud['X']['2019']['sales'], 1)             # het origineel blijft ongemoeid

    def test_tweede_bron_vergelijkt_alleen_eigen_cijfers(self):
        g = {'jaren': [{'year': 2020, 'comparable_eps': 2.0, 'sales': 1000.0, 'net_income': 100.0},
                       {'year': 2021, 'comparable_eps': 2.5, 'sales': 1100.0, 'equity': 50.0, 'aangeleverd': ['equity']}]}
        B.vergelijk_tweede_bron(g, 'X', {'bron': 'EODHD', 'fondsen': {'X': {'2020': {'comparable_eps': 2.04, 'sales': 1000000.0, 'net_income': 120.0},
                                                                            '2021': {'comparable_eps': 2.5, 'equity': 999}}}})
        tb = g['tweede_bron']
        self.assertEqual((tb['vergeleken'], tb['overeen']), (4, 2))   # eps 2020 (binnen 3%) en eps 2021 kloppen; aangeleverd eigen vermogen telt niet
        soorten = {x['veld']: x['soort'] for x in tb['afwijkend']}
        self.assertEqual(soorten, {'sales': 'eenheid', 'net_income': None})
        B.vergelijk_tweede_bron(g, 'X', {})
        self.assertNotIn('tweede_bron', g)

    def test_sjabloon_meldt_afgewezen_cijfers_valuta_en_heeft_controle(self):
        import json as _json
        sys.path.insert(0, str(Path(__file__).resolve().parent / 'hulp'))
        import aanvulsjabloon as S
        data = {'gemaakt': '2026-10-27T00:00:00Z', 'aandelen': [{'ticker': 'TTE.PA', 'naam': 'Total', 'valuta': 'EUR', 'gegevens': {
            'currencies': ['EUR', 'USD'],
            'jaren': [{'year': y, 'end': f'{y}-12-31', 'currency': 'USD', 'comparable_eps': 5.0, 'sales': 2e11} for y in range(2021, 2026)] +
                     [{'year': 2020, 'end': '2020-12-31', 'currency': 'USD', 'net_income': 7.0, 'aangeleverd': ['net_income'], 'aangeleverd_rij': True}],
            'aangeleverd': {'afgewezen': [{'jaar': 2019, 'veld': 'sales', 'waarde': 5e14, 'buur': 2e11}], 'valuta_jaren': [{'jaar': 2018, 'valuta': 'EUR', 'nodig': 'USD'}]}}}]}
        d = Path(tempfile.mkdtemp()); inp = d / 'a.json'; inp.write_text(_json.dumps(data)); uit = d / 'u.xlsx'
        S.main(str(inp), str(uit))
        wb = openpyxl.load_workbook(uit)
        ws = wb['Aanvullen']
        rij = {int(r[2].value): r for r in ws.iter_rows(min_row=3)}
        self.assertEqual(rij[2018][4].value, 'USD')                                  # rapportagevaluta, niet die van de beurs
        self.assertIn('stond in EUR', rij[2018][4].comment.text)
        kop = [c.value for c in ws[2]]
        self.assertIn('factor 10', rij[2019][kop.index('sales')].comment.text)
        self.assertIn('door u aangeleverd', rij[2020][kop.index('net_income')].comment.text)
        k = wb['Controle']
        jaren = [int(r[2].value) for r in k.iter_rows(min_row=3)]
        self.assertEqual(jaren, [2021, 2022, 2023, 2024, 2025])                      # het aangeleverde jaar 2020 niet
        # teruggestuurd met een tweede bron
        kk = [c.value for c in k[2]]
        k.cell(row=3, column=kk.index('comparable_eps') + 1, value=5.1)
        wb.save(uit)
        ctl, log = AI.lees_controle(uit)
        self.assertEqual(ctl, {'TTE.PA': {'2021': {'comparable_eps': 5.1}}})


class ValutaWissel(unittest.TestCase):
    def test_bestaand_jaar_in_eigen_valuta(self):
        g = {'jaren': [{'year': 2021, 'currency': 'GBP', 'sales': 100.0}, {'year': 2024, 'currency': 'USD', 'sales': 130.0}]}
        B.vul_aangeleverd(g, 'X', reg({2021: {'net_income': 10.0, 'currency': 'GBP'}, 2020: {'sales': 95.0, 'currency': 'GBP'}}))
        j = {x['year']: x for x in g['jaren']}
        self.assertEqual(j[2021]['net_income'], 10.0)                     # bestaand GBP-jaar: GBP is goed
        self.assertNotIn(2020, j)                                         # nieuw jaar: rapportagevaluta (USD) nodig
        self.assertEqual(g['aangeleverd']['valuta_jaren'], [{'jaar': 2020, 'valuta': 'GBP', 'nodig': 'USD'}])


class Correctie(unittest.TestCase):
    def tb(self, jaren):
        return {'bron': 'EODHD', 'fondsen': {'X': {str(y): r for y, r in jaren.items()}}}

    def test_splitsing_aena(self):
        g = {'niveau': 'jaarverslag', 'jaren': [
            {'year': 2022, 'comparable_eps': 6.01, 'net_income': 906e6, 'shares_diluted': 150e6},
            {'year': 2023, 'comparable_eps': 10.87, 'net_income': 1645e6, 'shares_diluted': 150e6},
            {'year': 2024, 'comparable_eps': 12.89, 'net_income': 1972e6, 'shares_diluted': 150e6},
            {'year': 2025, 'comparable_eps': 1.42, 'net_income': 2185e6, 'shares_diluted': 1500e6}]}
        reg = self.tb({2022: {'comparable_eps': 0.601, 'net_income': 901e6}, 2023: {'comparable_eps': 1.0872, 'net_income': 1631e6},
                       2024: {'comparable_eps': 1.2895, 'net_income': 1934e6}, 2025: {'comparable_eps': 1.4245, 'net_income': 2137e6}})
        B.vergelijk_tweede_bron(g, 'X', reg)
        eps = [j['comparable_eps'] for j in g['jaren']]
        self.assertEqual(eps, [0.601, 1.0872, 1.2895, 1.42])
        self.assertEqual([c['regel'] for c in g['tweede_bron']['gecorrigeerd']], ['splitsing'] * 3)
        self.assertEqual(g['niveau'], 'indicatief')
        B.vergelijk_tweede_bron(g, 'X', reg)                                   # opnieuw: zelfde uitkomst
        self.assertEqual([j['comparable_eps'] for j in g['jaren']], eps)
        B.vergelijk_tweede_bron(g, 'X', {})                                    # zonder tweede bron: alles terug
        self.assertEqual(g['jaren'][0]['comparable_eps'], 6.01); self.assertEqual(g['niveau'], 'jaarverslag')

    def test_shell_tweede_bron_fout_blijft_staan(self):
        g = {'jaren': [{'year': 2020 + i, 'comparable_eps': e, 'net_income': n, 'shares_diluted': s} for i, (e, n, s) in
                       enumerate([(2.57, 20.6e9, 8.0e9), (2.0, 16e9, 7.9e9), (5.71, 42.9e9, 7.5e9), (2.85, 19.6e9, 6.9e9)])]}
        reg = self.tb({2020: {'comparable_eps': 12.85, 'net_income': 20.1e9}, 2021: {'comparable_eps': 10.0, 'net_income': 15.8e9},
                       2022: {'comparable_eps': 5.709, 'net_income': 42.3e9}, 2023: {'comparable_eps': 2.847, 'net_income': 19.4e9}})
        B.vergelijk_tweede_bron(g, 'X', reg)
        self.assertEqual(g['jaren'][0]['comparable_eps'], 2.57)
        self.assertEqual(g['tweede_bron']['gecorrigeerd'], [])

    def test_narekenen_cellnex_en_bevestigd_nooit(self):
        rijen = [{'year': 2021, 'comparable_eps': -0.358, 'net_income': -386e6, 'shares_diluted': 700e6},
                 {'year': 2022, 'comparable_eps': -0.45, 'net_income': -313e6, 'shares_diluted': 705e6}]
        reg = self.tb({2021: {'comparable_eps': -0.5938}, 2022: {'comparable_eps': -0.4415}})
        g = {'jaren': copy.deepcopy(rijen)}
        B.vergelijk_tweede_bron(g, 'X', reg)
        self.assertEqual(g['jaren'][0]['comparable_eps'], -0.5938)
        self.assertEqual(g['tweede_bron']['gecorrigeerd'][0]['regel'], 'narekenen')
        b = {'eps_status': 'independently_verified', 'eps_verified_years': 10, 'jaren': copy.deepcopy(rijen)}
        B.vergelijk_tweede_bron(b, 'X', reg)
        self.assertEqual(b['jaren'][0]['comparable_eps'], -0.358)

    def test_aangeleverde_cijfers_tellen_niet_als_bewijs(self):
        # Merck 2016: nettowinst en aandelen kwamen uit de tweede bron zelf; dan geen correctie
        g = {'jaren': [{'year': 2016, 'comparable_eps': 3.75, 'net_income': 3920e6, 'shares_diluted': 434e6, 'aangeleverd': ['net_income', 'shares_diluted']},
                       {'year': 2017, 'comparable_eps': 5.99, 'net_income': 2605e6, 'shares_diluted': 435e6}]}
        B.vergelijk_tweede_bron(g, 'X', self.tb({2016: {'comparable_eps': 9.02}, 2017: {'comparable_eps': 5.99}}))
        self.assertEqual(g['jaren'][0]['comparable_eps'], 3.75)

    def test_aena_echt_met_uitschieter_en_aangeleverde_aandelen(self):
        g = {'niveau': 'jaarverslag', 'jaren': [
            {'year': 2021, 'comparable_eps': -3.17, 'net_income': -505e6, 'shares_diluted': 222e6},
            {'year': 2022, 'comparable_eps': 6.01, 'net_income': 906e6, 'shares_diluted': 1500e6, 'aangeleverd': ['shares_diluted']},
            {'year': 2023, 'comparable_eps': 10.87, 'net_income': 1645e6, 'shares_diluted': 150e6, 'aangeleverd': ['shares_diluted']},
            {'year': 2024, 'comparable_eps': 12.89, 'net_income': 1972e6, 'shares_diluted': 150e6, 'aangeleverd': ['shares_diluted']},
            {'year': 2025, 'comparable_eps': 1.42, 'net_income': 2185e6, 'shares_diluted': 1500e6}]}
        reg = self.tb({2021: {'comparable_eps': -0.04, 'net_income': -60e6}, 2022: {'comparable_eps': 0.601, 'net_income': 901e6},
                       2023: {'comparable_eps': 1.0872, 'net_income': 1631e6}, 2024: {'comparable_eps': 1.2895, 'net_income': 1934e6},
                       2025: {'comparable_eps': 1.4245, 'net_income': 2137e6}})
        B.vergelijk_tweede_bron(g, 'X', reg)
        self.assertEqual([j['comparable_eps'] for j in g['jaren']], [-3.17, 0.601, 1.0872, 1.2895, 1.42])


class Valutawissel(unittest.TestCase):
    def g(self):
        return {'niveau': 'jaarverslag', 'currencies': ['GBP', 'USD'], 'jaren': [
            {'year': 2020, 'currency': 'GBP', 'comparable_eps': 1.0, 'sales': 10e9, 'net_income': 2e9, 'shares_diluted': 2e9},
            {'year': 2021, 'currency': 'GBP', 'comparable_eps': 1.1, 'sales': 11e9, 'net_income': 2.2e9, 'shares_diluted': 2e9, 'koers_hoog': 30},
            {'year': 2022, 'currency': 'USD', 'comparable_eps': 1.5, 'sales': 15e9, 'net_income': 3e9},
            {'year': 2023, 'currency': 'USD', 'comparable_eps': 1.6, 'sales': 16e9, 'net_income': 3.2e9}]}

    def test_oud_jaar_in_nieuwe_valuta_vervangt_heel_jaar_en_terug(self):
        g = self.g()
        B.vul_aangeleverd(g, 'X', reg({2021: {'currency': 'USD', 'comparable_eps': 1.4, 'sales': 14e9}}))
        j = {x['year']: x for x in g['jaren']}
        self.assertEqual((j[2021]['currency'], j[2021]['comparable_eps'], j[2021]['sales']), ('USD', 1.4, 14e9))
        self.assertNotIn('net_income', j[2021])                   # geen GBP-bedrag naast USD-bedragen
        self.assertEqual(j[2021]['shares_diluted'], 2e9); self.assertEqual(j[2021]['koers_hoog'], 30)
        self.assertEqual(j[2020]['currency'], 'GBP')              # niet aangeleverd: blijft
        self.assertEqual(g['aangeleverd']['valuta_vervangen'], [{'jaar': 2021, 'was': 'GBP', 'nu': 'USD'}])
        self.assertEqual(g['niveau'], 'indicatief')
        B.vul_aangeleverd(g, 'X', {'fondsen': {}})
        j = {x['year']: x for x in g['jaren']}
        self.assertEqual((j[2021]['currency'], j[2021]['comparable_eps'], j[2021]['net_income']), ('GBP', 1.1, 2.2e9))
        self.assertEqual(g['currencies'], ['GBP', 'USD']); self.assertEqual(g['niveau'], 'jaarverslag')

    def test_sjabloon_vraagt_oude_valutajaren_opnieuw(self):
        import json as _json
        import aanvulsjabloon as S
        data = {'gemaakt': '2026-10-27T00:00:00Z', 'aandelen': [{'ticker': 'DGE.L', 'naam': 'Diageo', 'valuta': 'GBP', 'gegevens': self.g()}]}
        d = Path(tempfile.mkdtemp()); inp = d / 'a.json'; inp.write_text(_json.dumps(data)); uit = d / 'u.xlsx'
        S.main(str(inp), str(uit))
        ws = openpyxl.load_workbook(uit)['Aanvullen']
        kop = [c.value for c in ws[2]]
        rij = {int(r[2].value): r for r in ws.iter_rows(min_row=3)}
        self.assertEqual(rij[2021][4].value, 'USD')
        self.assertIn('nog in GBP', rij[2021][4].comment.text)
        self.assertIsNone(rij[2021][kop.index('comparable_eps')].value)
        self.assertIn('1.1 GBP', rij[2021][kop.index('comparable_eps')].comment.text)
        self.assertEqual(rij[2021][kop.index('shares_diluted')].value, 2e9)
        self.assertEqual(rij[2022][kop.index('comparable_eps')].value, 1.5)


class KwartaalWinst(unittest.TestCase):
    def g(self):
        return {'jaren': [{'year': 2025, 'end': '2025-12-31', 'currency': 'EUR', 'comparable_eps': 3.0}]}

    def ts(self, d):
        import datetime as dt
        return int(dt.datetime.fromisoformat(d).timestamp())

    def test_gebruikt_alleen_recent_geloofwaardig_zelfde_valuta(self):
        info = {'trailingEps': 3.4, 'financialCurrency': 'EUR', 'mostRecentQuarter': self.ts('2026-09-30')}
        t = B.ttm_uit_info(info, self.g(), '2026-10-27')
        self.assertEqual((t['eps'], t['tot']), (3.4, '2026-09-30'))
        self.assertIsNone(B.ttm_uit_info(dict(info, financialCurrency='USD'), self.g(), '2026-10-27'))
        self.assertIsNone(B.ttm_uit_info(dict(info, mostRecentQuarter=self.ts('2025-12-31')), self.g(), '2026-10-27'))   # = jaarwinst
        self.assertIsNone(B.ttm_uit_info(dict(info, trailingEps=6.6), self.g(), '2026-10-27'))                           # factor 2,2
        self.assertEqual(B.ttm_uit_info(dict(info, trailingEps=340.0), self.g(), '2026-10-27')['eps'], 3.4)               # pence
        self.assertIsNone(B.ttm_uit_info(None, self.g(), '2026-10-27'))


class TweeBronnen(unittest.TestCase):
    def g(self, n=6):
        return {'niveau': 'indicatief', 'bron': 'Yahoo Finance', 'jaren': [{'year': 2020 + i, 'comparable_eps': 1 + i / 10, 'eps_bron': 'yahoo', 'sales': 100.0} for i in range(n)]}

    def test_eens_bij_genoeg_overeenkomende_jaren(self):
        g = self.g()
        reg = {'bron': 'EODHD', 'fondsen': {'X': {str(2020 + i): {'comparable_eps': 1 + i / 10} for i in range(6)}}}
        B.vergelijk_tweede_bron(g, 'X', reg)
        self.assertEqual(g['twee_bronnen_eens']['jaren'], 6)
        reg['fondsen']['X']['2023']['comparable_eps'] = 2.0                          # één afwijkend jaar: niet eens
        B.vergelijk_tweede_bron(g, 'X', reg)
        self.assertNotIn('twee_bronnen_eens', g)
        g2 = self.g(); g2['jaren'][0]['aangeleverd'] = ['comparable_eps']             # winst uit de tweede bron zelf: telt niet
        B.vergelijk_tweede_bron(g2, 'X', {'bron': 'EODHD', 'fondsen': {'X': {str(2020 + i): {'comparable_eps': 1 + i / 10} for i in range(6)}}})
        self.assertNotIn('twee_bronnen_eens', g2)


class Adr(unittest.TestCase):
    def test_shell_per_ads_en_terug(self):
        g = {'jaren': [{'year': 2025, 'comparable_eps': 3.0, 'reported_eps': 3.0, 'net_income': 18e9, 'shares_diluted': 6e9}]}
        B.adr_toe(g, 'SHEL')
        j = g['jaren'][0]
        self.assertEqual((j['comparable_eps'], j['shares_diluted'], j['net_income']), (6.0, 3e9, 18e9))
        B.adr_toe(g, 'SHEL')                                   # niet dubbel
        self.assertEqual(g['jaren'][0]['comparable_eps'], 6.0)
        B.adr_terug(g)
        self.assertEqual((g['jaren'][0]['comparable_eps'], g['jaren'][0]['shares_diluted']), (3.0, 6e9)); self.assertNotIn('adr', g)
        a = {'jaren': [{'year': 2025, 'comparable_eps': 6.54}]}
        B.adr_toe(a, 'AZN'); self.assertEqual(a['jaren'][0]['comparable_eps'], 3.27)
        n = {'jaren': [{'year': 2025, 'comparable_eps': 7.15}]}
        B.adr_toe(n, 'NVS'); self.assertEqual(n['jaren'][0]['comparable_eps'], 7.15)


class PdfNarekenen(unittest.TestCase):
    def g(self, n=10):
        return {'niveau': 'indicatief', 'bron': 'Yahoo Finance', 'currencies': ['EUR'], 'jaren': [
            {'year': 2016 + i, 'currency': 'EUR', 'comparable_eps': 2.0 + i / 10, 'eps_bron': 'jaarverslag', 'net_income': (2.0 + i / 10) * 400e6} for i in range(n)]}

    def reg(self, n=10, fout=None):
        return {'X': {str(2016 + i): {'document': 'u', 'kandidaten': [{'aandelen': 999e6}, {'aandelen': 401e6 if i != fout else 300e6}]} for i in range(n)}}

    def test_tien_jaar_nagerekend_is_bevestigd_en_terug(self):
        g = self.g()
        B.reken_pdf_eps_na(g, 'X', self.reg())
        self.assertTrue(g['pdf_narekening']['bevestigd']); self.assertEqual(g['eps_status'], 'independently_verified')
        self.assertEqual(g['jaren'][0]['eps_narekening']['aandelen'], 401e6)
        B.reken_pdf_eps_na(g, 'X', {})
        self.assertNotIn('eps_status', g); self.assertNotIn('eps_narekening', g['jaren'][0])

    def test_een_jaar_klopt_niet_dan_niet_bevestigd(self):
        g = self.g()
        B.reken_pdf_eps_na(g, 'X', self.reg(fout=5))
        self.assertEqual(g['pdf_narekening']['jaren'], 9); self.assertNotIn('bevestigd', g['pdf_narekening'])
        self.assertNotIn('eps_status', g)
        a = self.g(); a['jaren'][0]['aangeleverd'] = ['net_income']        # aangeleverde nettowinst telt wel (onafhankelijke derde bron)
        B.reken_pdf_eps_na(a, 'X', self.reg())
        self.assertEqual(a['pdf_narekening']['jaren'], 10)
