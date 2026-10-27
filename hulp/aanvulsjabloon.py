"""Hulpmiddel: Excel om ontbrekende jaarcijfers aan te vullen, vooraf gevuld met wat de app al heeft.
Gebruik: python hulp/aanvulsjabloon.py web/data/aandelen.json uit.xlsx

Per fonds met een gat in de laatste 10 boekjaren: één rij per boekjaar. Bekende cijfers staan er grijs in (met de bron),
ontbrekende cellen zijn geel. Zelfde kolommen als het tabblad 'Sjabloon jaarcijfers' van de specificatie."""
import json, sys
from datetime import date, timedelta

VELDEN = [  # (kolom, kop, prioriteit)
    ('reported_eps', 'Winst/aandeel gerapporteerd', 1), ('comparable_eps', 'Winst/aandeel vergelijkbaar', 1),
    ('sales', 'Omzet', 1), ('pretax_income', 'Winst vóór belasting', 1), ('net_income', 'Nettowinst aandeelhouders', 1),
    ('equity', 'Eigen vermogen aandeelhouders', 1), ('shares_diluted', 'Gem. aantal aandelen (verwaterd)', 2),
    ('operating_cash_flow', 'Kasstroom bedrijfsactiviteiten', 2), ('capex', 'Investeringen vaste activa', 2),
    ('total_assets', 'Totaal bezittingen', 2), ('total_liabilities', 'Totaal verplichtingen', 2),
]
JAREN = 10


def getal(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) and x == x else None


def bron_van(g, j, veld):
    if veld in (j.get('aangeleverd') or []):
        return 'door u aangeleverd'
    if veld in (j.get('jaarcijfers_pdf') or []):
        return 'pdf-jaarverslag'
    if veld in ('reported_eps', 'comparable_eps') and j.get('eps_bron'):
        return {'jaarverslag': 'pdf-jaarverslag', 'yahoo': 'Yahoo'}.get(j['eps_bron'], j['eps_bron'])
    return {'Yahoo Finance': 'Yahoo'}.get(g.get('bron'), g.get('bron') or '')


def rijen_voor(rec, peildatum=None):
    g = rec.get('gegevens') or {}
    if g.get('fout') and not g.get('jaren'):
        g = {}
    jaren = sorted((j for j in g.get('jaren') or [] if isinstance(j.get('year'), int)), key=lambda j: j['year'])
    per = {j['year']: j for j in jaren}
    if jaren:
        laatste = jaren[-1]
        try:
            e = date.fromisoformat(laatste['end'])
            md = (e.month, e.day)
        except (KeyError, ValueError, TypeError):
            md = (12, 31)
        tot = laatste['year']
    else:
        md, tot = (12, 31), 0
    # het nieuwste boekjaar waarvan het jaarverslag er nu zou moeten zijn (einddatum minstens 4 maanden voor de gegevensdatum)
    ref = (peildatum or date.today()) - timedelta(days=120)
    verwacht = ref.year if md <= (ref.month, ref.day) else ref.year - 1
    tot = max(tot, verwacht)
    uit = []
    for y in range(tot - JAREN + 1, tot + 1):
        j = per.get(y) or {}
        try:
            eind = j.get('end') or date(y, md[0], min(md[1], 28 if md[0] == 2 else md[1])).isoformat()
        except ValueError:
            eind = f'{y}-12-31'
        waarden = {v: getal(j.get(v)) for v, _, _ in VELDEN}
        r_adr = (g.get('adr') or {}).get('aandelen_per_ads')
        if r_adr:                                     # het sjabloon werkt per gewoon aandeel, zoals de jaarverslagen
            for v in ('reported_eps', 'comparable_eps'):
                if waarden[v] is not None:
                    waarden[v] = round(waarden[v] / r_adr, 6)
            if waarden['shares_diluted'] is not None:
                waarden['shares_diluted'] = round(waarden['shares_diluted'] * r_adr)
        oud = None
        munt = jaren[-1].get('currency') if jaren else None
        if j.get('currency') and munt and j['currency'] != munt:
            # het bedrijf rapporteert nu in een andere valuta: dit jaar opnieuw, in de huidige valuta (aandelen blijven)
            oud = {'valuta': j['currency'], 'waarden': {v: x for v, x in waarden.items() if x is not None and v != 'shares_diluted'}}
            waarden = {v: (x if v == 'shares_diluted' else None) for v, x in waarden.items()}
            j = dict(j, currency=munt)
        bronnen = {v: bron_van(g, j, v) for v, _, _ in VELDEN if waarden[v] is not None}
        # valuta: die van het jaar zelf, anders die van het laatste boekjaar (de rapportagevaluta), pas daarna die van de beurs
        uit.append({'year': y, 'end': eind, 'currency': j.get('currency') or (jaren[-1].get('currency') if jaren else None) or (g.get('currencies') or [rec.get('valuta')])[0],
                    'waarden': waarden, 'bronnen': bronnen, 'oud': oud})
    return uit, g


def main(inpad, uitpad):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    data = json.load(open(inpad, encoding='utf-8'))
    F = 'Arial'
    KOP = Font(name=F, bold=True, color='FFFFFF'); KOPFILL = PatternFill('solid', fgColor='1F4E5A')
    GEEL = PatternFill('solid', fgColor='FFFF00'); GRIJS = PatternFill('solid', fgColor='E7E6E6')
    GEWOON = Font(name=F, size=10); VET = Font(name=F, size=10, bold=True)
    fondsen = []
    try:
        peil = date.fromisoformat(str(data.get('gemaakt') or '')[:10])
    except ValueError:
        peil = date.today()
    for rec in data.get('aandelen') or []:
        rijen, g = rijen_voor(rec, peil)
        ontbreekt1 = sum(1 for r in rijen for v, _, p in VELDEN if p == 1 and r['waarden'][v] is None)
        ontbreekt2 = sum(1 for r in rijen for v, _, p in VELDEN if p == 2 and r['waarden'][v] is None)
        if ontbreekt1 + ontbreekt2 == 0:
            continue
        eps_jaren = sum(1 for r in rijen if r['waarden']['comparable_eps'] is not None)
        fondsen.append((rec, g, rijen, ontbreekt1, ontbreekt2, eps_jaren))
    # de grootste gaten in de winst per aandeel eerst (die blokkeren het NAIC-oordeel), dan de overige gaten
    fondsen.sort(key=lambda x: (x[5] >= 4, x[5], -x[3], x[0].get('ticker')))

    wb = Workbook()
    ws = wb.active; ws.title = 'Leeswijzer'
    tekst = [
        ('Aandeel Signaal – ontbrekende jaarcijfers aanvullen', Font(name=F, size=14, bold=True)),
        (f"Gemaakt uit de gegevens van {str(data.get('gemaakt') or '')[:10]}. {len(fondsen)} fondsen met minstens één ontbrekend cijfer in de laatste {JAREN} boekjaren.", GEWOON),
        ('', GEWOON),
        ('Zo vult u aan', VET),
        ('Tabblad Aanvullen: één rij per fonds per boekjaar. Gele cellen ontbreken nog; vul daar het cijfer in. Grijze cellen heeft de app al (de bron staat in de opmerking bij de cel); laat die staan.', GEWOON),
        ('Ziet u dat een grijs cijfer niet klopt, zet dan het juiste cijfer in de kolom Opmerking van die rij, met de bron. De app vervangt een officieel cijfer nooit automatisch.', GEWOON),
        ('Bedragen in hele valuta (9.850.000.000, niet 9.850 of 9,85 mld). Winst per aandeel in valuta per aandeel (niet in centen of pence). Leeg = onbekend, nooit 0 voor een ontbrekend cijfer.', GEWOON),
        ('Winst per aandeel vergelijkbaar: na splitsingen, samenvoegingen of bonusaandelen omgerekend naar het huidige aantal aandelen. Zonder zulke gebeurtenissen gelijk aan de gerapporteerde.', GEWOON),
        ('Nettowinst, winst per aandeel en eigen vermogen: deel van de aandeelhouders van het moederbedrijf.', GEWOON),
        ('Tabblad Herkomst (sterk aanbevolen): per ingevuld cijfer of per jaarverslag het document en de pagina. Met herkomst kan de app de cijfers als “uit jaarverslag” tonen; zonder blijven ze “aangeleverd, niet gecontroleerd”.', GEWOON),
        ('Kolom Controle EPS rekent nettowinst ÷ aantal aandelen; die hoort dicht bij de gerapporteerde winst per aandeel te liggen.', GEWOON),
        ('', GEWOON),
        ('Prioriteit', VET),
        ('Tabblad Overzicht staat op volgorde: eerst fondsen met minder dan 4 jaar winst per aandeel (zonder die cijfers geen NAIC-oordeel), dan de rest. Prioriteit 1-velden zijn het belangrijkst voor het NAIC-oordeel.', GEWOON),
        ('Aangeleverde cijfers zijn zichtbaar voor alle gebruikers van de app. Gebruik alleen bronnen die u mag publiceren.', GEWOON),
        ('', GEWOON),
        ('Opmerkingen bij gele cellen', VET),
        ('Staat er een opmerking bij een gele cel of bij de valuta, dan heeft u dat cijfer al eens aangeleverd maar is het niet gebruikt: in een andere valuta dan waarin het bedrijf rapporteert, of meer dan een factor 10 anders dan het jaar ernaast. De opmerking zegt welk getal en waarom. Eerder aangeleverde cijfers die wel gebruikt zijn, staan grijs met als bron “door u aangeleverd”; u hoeft ze niet opnieuw te sturen.', GEWOON),
        ('', GEWOON),
        ('Tabblad Controle (tweede bron)', VET),
        ('Hier staan de cijfers die de app uit officiële bronnen heeft (SEC, ESEF, jaarverslag of Yahoo). Vul in de gele kolom “uw bron” hetzelfde cijfer uit uw eigen bron in. De app vervangt niets, maar vergelijkt: afwijkingen staan daarna bij het fonds onder Bronnen. Zo komen fouten aan het licht, zoals een gemiste splitsing of een verkeerde eenheid.', GEWOON),
    ]
    for i, (t, f) in enumerate(tekst, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120

    ov = wb.create_sheet('Overzicht')
    kop = ['Ticker', 'Naam', 'Bron nu', 'Jaren met winst/aandeel', 'Ontbrekend prioriteit 1', 'Ontbrekend prioriteit 2', 'Boekjaren']
    for c, k in enumerate(kop, 1):
        cel = ov.cell(row=1, column=c, value=k); cel.font = KOP; cel.fill = KOPFILL; cel.alignment = Alignment(wrap_text=True)
    for r, (rec, g, rijen, o1, o2, ej) in enumerate(fondsen, 2):
        for c, v in enumerate([rec.get('ticker'), rec.get('naam'), {'Yahoo Finance': 'Yahoo (indicatief)'}.get(g.get('bron'), g.get('bron') or 'geen jaarcijfers'),
                               ej, o1, o2, f"{rijen[0]['year']}–{rijen[-1]['year']}"], 1):
            ov.cell(row=r, column=c, value=v).font = GEWOON
    for c, b in enumerate([12, 40, 20, 14, 14, 14, 14], 1):
        ov.column_dimensions[get_column_letter(c)].width = b
    ov.freeze_panes = 'C2'; ov.auto_filter.ref = f'A1:G{len(fondsen) + 1}'

    a = wb.create_sheet('Aanvullen')
    kop = ['ticker', 'naam', 'year', 'end', 'currency'] + [v for v, _, _ in VELDEN] + ['controle_eps', 'Opmerking']
    namen = ['Ticker', 'Naam', 'Boekjaar', 'Einddatum', 'Valuta'] + [f'{k} (P{p})' for _, k, p in VELDEN] + ['Controle EPS (nettowinst ÷ aandelen)', 'Opmerking / correctie met bron']
    for c, (k, n) in enumerate(zip(kop, namen), 1):
        cel = a.cell(row=1, column=c, value=n); cel.font = KOP; cel.fill = KOPFILL; cel.alignment = Alignment(wrap_text=True, vertical='top')
        cel2 = a.cell(row=2, column=c, value=k); cel2.font = Font(name=F, size=8, italic=True, color='808080')
    kol = {k: i + 1 for i, k in enumerate(kop)}
    r = 3
    for rec, g, rijen, o1, o2, ej in fondsen:
        for rij in rijen:
            a.cell(row=r, column=1, value=rec.get('ticker')).font = GEWOON
            a.cell(row=r, column=2, value=rec.get('naam')).font = GEWOON
            a.cell(row=r, column=3, value=str(rij['year'])).font = GEWOON
            a.cell(row=r, column=4, value=rij['end']).font = GEWOON
            a.cell(row=r, column=5, value=rij['currency']).font = GEWOON
            av = g.get('aangeleverd') or {}
            afg = {(x['jaar'], x['veld']): x for x in av.get('afgewezen') or []}
            vj = {x['jaar']: x for x in av.get('valuta_jaren') or []}
            if rij['year'] in vj and vj[rij['year']]['valuta'] != rij['currency']:
                x = vj[rij['year']]
                a.cell(row=r, column=5).comment = Comment(f"Vorige aanlevering stond in {x['valuta']} en is niet gebruikt: dit bedrijf rapporteert in {x['nodig']}. Vul de cijfers in {x['nodig']} in.", 'Aandeel Signaal')
                a.cell(row=r, column=5).font = VET
            if rij.get('oud'):
                a.cell(row=r, column=5).comment = Comment(f"Dit jaar staat in de app nog in {rij['oud']['valuta']}, terwijl het bedrijf nu in {rij['currency']} rapporteert. Vul alle cijfers van dit jaar in {rij['currency']} in (bijv. het vergelijkende cijfer in het jaarverslag van na de valutawissel); de app vervangt dan het hele jaar.", 'Aandeel Signaal')
                a.cell(row=r, column=5).font = VET
            for v, _, _ in VELDEN:
                cel = a.cell(row=r, column=kol[v], value=rij['waarden'][v]); cel.font = GEWOON
                if rij['waarden'][v] is None:
                    cel.fill = GEEL
                    x = afg.get((rij['year'], v))
                    if rij.get('oud') and v in rij['oud']['waarden']:
                        cel.comment = Comment(f"Nu in de app: {rij['oud']['waarden'][v]:,} {rij['oud']['valuta']}. Vul het cijfer in {rij['currency']} in.", 'Aandeel Signaal')
                    elif x:
                        cel.comment = Comment(f"Vorige aanlevering: {x['waarde']:,} niet gebruikt; wijkt meer dan een factor 10 af van {x['buur']:,} in een jaar ernaast (andere eenheid of ander begrip?). Controleer en vul opnieuw in, of laat leeg.", 'Aandeel Signaal')
                else:
                    cel.fill = GRIJS
                    if rij['bronnen'].get(v):
                        cel.comment = Comment(f"Bron: {rij['bronnen'][v]}", 'Aandeel Signaal')
                cel.number_format = '0.00' if v.endswith('eps') else '#,##0;(#,##0);-'
            ni, sh = get_column_letter(kol['net_income']), get_column_letter(kol['shares_diluted'])
            c = a.cell(row=r, column=kol['controle_eps'], value=f'=IF(AND(ISNUMBER({ni}{r}),ISNUMBER({sh}{r}),{sh}{r}<>0),{ni}{r}/{sh}{r},"")')
            c.font = GEWOON; c.number_format = '0.00'
            a.cell(row=r, column=kol['Opmerking']).font = GEWOON
            r += 1
    breedte = {1: 10, 2: 30, 3: 9, 4: 12, 5: 8}
    for c in range(1, len(kop) + 1):
        a.column_dimensions[get_column_letter(c)].width = breedte.get(c, 16)
    a.column_dimensions[get_column_letter(len(kop))].width = 40
    a.row_dimensions[1].height = 45
    a.freeze_panes = 'F3'; a.auto_filter.ref = f'A2:{get_column_letter(len(kop))}{r - 1}'

    # Tabblad Controle: de cijfers van de app (officieel, jaarverslag of Yahoo; niet de aangeleverde) naast een lege kolom voor
    # dezelfde cijfers uit een tweede bron. Teruggestuurd vergelijkt de app ze en toont de verschillen bij Bronnen.
    CV = [('comparable_eps', 'Winst/aandeel vergelijkbaar'), ('sales', 'Omzet'), ('net_income', 'Nettowinst aandeelhouders'), ('equity', 'Eigen vermogen aandeelhouders')]
    k = wb.create_sheet('Controle')
    ckop = ['ticker', 'naam', 'year', 'end', 'currency'] + [x for v, _ in CV for x in (v + '_app', v)]
    cnaam = ['Ticker', 'Naam', 'Boekjaar', 'Einddatum', 'Valuta'] + [x for _, n in CV for x in (n + ' – app', n + ' – uw bron')]
    for c, (kk, n) in enumerate(zip(ckop, cnaam), 1):
        cel = k.cell(row=1, column=c, value=n); cel.font = KOP; cel.fill = KOPFILL; cel.alignment = Alignment(wrap_text=True, vertical='top')
        k.cell(row=2, column=c, value=kk).font = Font(name=F, size=8, italic=True, color='808080')
    rk = 3
    for rec in data.get('aandelen') or []:
        g = rec.get('gegevens') or {}
        for j in sorted((j for j in g.get('jaren') or [] if isinstance(j.get('year'), int)), key=lambda j: j['year'])[-JAREN:]:
            eigen = {v: getal(j.get(v)) for v, _ in CV if v not in (j.get('aangeleverd') or [])}
            if (g.get('adr') or {}).get('aandelen_per_ads') and eigen.get('comparable_eps') is not None:
                eigen['comparable_eps'] = round(eigen['comparable_eps'] / g['adr']['aandelen_per_ads'], 6)
            if not any(x is not None for x in eigen.values()):
                continue
            for c, v in enumerate([rec.get('ticker'), rec.get('naam'), str(j['year']), j.get('end'), j.get('currency')], 1):
                k.cell(row=rk, column=c, value=v).font = GEWOON
            for i, (v, _) in enumerate(CV):
                ca = k.cell(row=rk, column=6 + 2 * i, value=eigen.get(v)); ca.font = GEWOON; ca.fill = GRIJS
                cb = k.cell(row=rk, column=7 + 2 * i); cb.font = GEWOON
                if eigen.get(v) is not None:
                    cb.fill = GEEL
                for cc in (ca, cb):
                    cc.number_format = '0.00' if v.endswith('eps') else '#,##0;(#,##0);-'
            rk += 1
    for c in range(1, len(ckop) + 1):
        k.column_dimensions[get_column_letter(c)].width = {1: 10, 2: 30, 3: 9, 4: 12, 5: 8}.get(c, 16)
    k.row_dimensions[1].height = 45
    k.freeze_panes = 'F3'; k.auto_filter.ref = f'A2:{get_column_letter(len(ckop))}{rk - 1}'

    h = wb.create_sheet('Herkomst')
    kop = ['Ticker', 'Boekjaar', 'Veld (of: alle)', 'Bron (jaarverslag, website, …)', 'Document of url', 'Publicatiedatum (JJJJ-MM-DD)', 'Pagina', 'Opmerking']
    for c, k in enumerate(kop, 1):
        cel = h.cell(row=1, column=c, value=k); cel.font = KOP; cel.fill = KOPFILL; cel.alignment = Alignment(wrap_text=True)
        h.column_dimensions[get_column_letter(c)].width = [10, 10, 16, 26, 60, 18, 8, 40][c - 1]
    voorbeeld = ['ADS.DE', '2019', 'alle', 'jaarverslag', 'https://www.adidas-group.com/…/annual-report-2019.pdf', '2020-03-11', '', 'voorbeeld: verwijder deze rij']
    for c, v in enumerate(voorbeeld, 1):
        h.cell(row=2, column=c, value=v).font = Font(name=F, size=10, color='0000FF')
    for rr in range(3, 300):
        for c in range(1, len(kop) + 1):
            h.cell(row=rr, column=c).fill = GEEL
    h.freeze_panes = 'A2'
    wb.save(uitpad)
    print(len(fondsen), 'fondsen,', r - 3, 'rijen')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
